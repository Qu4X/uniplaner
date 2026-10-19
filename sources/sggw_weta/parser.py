"""
Parser siatki planów zajęć SGGW z arkuszy Excel (.xlsx).
"""
import re
import os
import logging
from datetime import datetime, timedelta
from typing import Dict, List, Optional, Tuple, Any, Union
import openpyxl

logger = logging.getLogger("sggw_parser")

DAY_NAME_MAP = {
    "poniedziałek": "PON",
    "poniedzialek": "PON",
    "poniedz.": "PON",
    "poniedz": "PON",
    "pon": "PON",
    "wtorek": "WT",
    "wt": "WT",
    "środa": "ŚR",
    "sroda": "ŚR",
    "śr": "ŚR",
    "sr": "ŚR",
    "czwartek": "CZW",
    "czw": "CZW",
    "piątek": "PT",
    "piatek": "PT",
    "pt": "PT"
}


def _normalize_time_24h(raw_time: str) -> str:
    clean = raw_time.replace(".", ":").strip()
    h_str, m_str = clean.split(":")
    h = int(h_str)
    if 1 <= h <= 6:
        h += 12
    return f"{h:02d}:{m_str}"


def _time_to_min(t: str) -> int:
    h, m = map(int, t.split(":"))
    return h * 60 + m


def parse_time_bounds(text: str, default_start: str, default_end: str) -> Tuple[str, str]:
    """
    Wyznacza godziny rozpoczęcia i zakończenia zajęć (w formacie 24h HH:MM).
    Jeśli w tekście podano jawne 'od XX:YY do AA:BB' lub 'od XX:YY' / 'do AA:BB', ma to priorytet.
    """
    start_time = _normalize_time_24h(default_start)
    end_time = _normalize_time_24h(default_end)

    # Format: od HH:MM do HH:MM lub HH:MM - HH:MM
    range_match = re.search(r'(?:od\s*)?(\d{1,2}[:.]\d{2})\s*(?:do|-)\s*(\d{1,2}[:.]\d{2})', text, re.IGNORECASE)
    if range_match:
        s = _normalize_time_24h(range_match.group(1))
        e = _normalize_time_24h(range_match.group(2))
        if _time_to_min(e) > _time_to_min(s):
            return (s, e)

    # Format: tylko "od HH:MM"
    from_match = re.search(r'\bod\s*(\d{1,2}[:.]\d{2})', text, re.IGNORECASE)
    if from_match:
        cand_s = _normalize_time_24h(from_match.group(1))
        if _time_to_min(cand_s) < _time_to_min(end_time):
            start_time = cand_s

    # Format: tylko "do HH:MM"
    to_match = re.search(r'\bdo\s*(\d{1,2}[:.]\d{2})', text, re.IGNORECASE)
    if to_match:
        cand_e = _normalize_time_24h(to_match.group(1))
        if _time_to_min(cand_e) > _time_to_min(start_time):
            end_time = cand_e

    if _time_to_min(end_time) <= _time_to_min(start_time):
        end_min = _time_to_min(start_time) + 90
        end_time = f"{end_min // 60:02d}:{end_min % 60:02d}"

    return (start_time, end_time)


CANONICAL_SUBJECT_MAP = {
    "anat. zw.": "Anatomia zwierząt",
    "anat. zw": "Anatomia zwierząt",
    "anat zw": "Anatomia zwierząt",
    "komunikowanie społ": "Komunikowanie społeczne",
    "komunikowanie społ.": "Komunikowanie społeczne",
    "mikrobiol": "Mikrobiologia",
    "admin weter": "Administracja weterynaryjna",
    "choroby wew. m. zw.": "Choroby wewnętrzne małych zwierząt",
    "choroby zakaźne m. zw.": "Choroby zakaźne małych zwierząt",
    "chirurgia m .zw.": "Chirurgia małych zwierząt",
    "chirurgia m. zw.": "Chirurgia małych zwierząt",
    "chirurgia małych zw": "Chirurgia małych zwierząt",
    "chirurgia małych zw.": "Chirurgia małych zwierząt",
    "bezpieczeństwo żywności poch. zw.": "Bezpieczeństwo żywności pochodzenia zwierzęcego",
    "rozród m. zw.": "Rozród małych zwierząt",
    "ch.zw.fut": "Choroby zwierząt futerkowych",
    "ch.zw. fut.": "Choroby zwierząt futerkowych",
    "chor. zw. fut.": "Choroby zwierząt futerkowych",
}


def extract_groups_from_text(text: str, all_groups: List[str]) -> List[str]:
    """
    Identyfikuje grupy ćwiczeniowe z tekstu komórki.
    Obsługuje zakresy (np. gr. 1-3, 4-7, 8-10), slashe (np. gr. 1/2, 4/7, 9/10),
    listy po przecinku i ze spójnikiem 'i' (np. gr 1 i 2, grupy 1,2,3,4).
    Liczby powyżej liczby grup (np. gr. 1,2– 1071) traktuje jako salę po myślniku.
    """
    clean_text = text.replace("\n", " ")
    max_group_num = len(all_groups) if all_groups else 10

    target_group_nums = set()

    # 1. Standardowe prefiksy grupowe: gr., gr, grupa, grupy, grup, (*)gr.
    group_matches = re.findall(
        r'(?:(?:\(\*?\)\s*)?\b(?:gr\.?|grupa|grupy|grup))\s*([0-9\s,–\-/i]+(?:\s+i\s+[0-9]+)?)',
        clean_text,
        re.IGNORECASE
    )

    # 2. Oznaczenie grupy bezpośrednio po kropce na końcu nazwy (np. "Choroby zakaźne małych zwierząt.3,4")
    dot_matches = re.findall(r'(?<=\w)\.([0-9\s,–\-/]+)(?:\s|$)', clean_text)
    if dot_matches:
        for dm in dot_matches:
            # Upewnij się, że to nie jest część daty ani godziny
            if not re.search(r'\d{4}|\d{2}:\d{2}', dm):
                group_matches.append(dm)

    for gm in group_matches:
        # A. Zakresy: 1-3, 4-7, 8-10, 1-10 (precyzyjnie 1-2 cyfry, nie prefiks większej liczby jak 1071)
        for r_m in re.finditer(r'(?<!\d)(\d{1,2})\s*[-–]\s*(\d{1,2})(?!\d)', gm):
            start_n, end_n = int(r_m.group(1)), int(r_m.group(2))
            if 1 <= start_n <= max_group_num and 1 <= end_n <= max_group_num and start_n <= end_n:
                for n in range(start_n, end_n + 1):
                    target_group_nums.add(n)

        # B. Slashe: 1/2, 4/7, 3/8, 5/6, 9/10
        for s_m in re.finditer(r'(?<!\d)(\d{1,2})\s*/\s*(\d{1,2})(?!\d)', gm):
            a, b = int(s_m.group(1)), int(s_m.group(2))
            if 1 <= a <= max_group_num:
                target_group_nums.add(a)
            if 1 <= b <= max_group_num:
                target_group_nums.add(b)

        # C. Pojedyncze liczby (np. "1, 2", "3,4,7", "1 i 2", "1,2,8", "gr 1")
        # Ignorujemy liczby większe niż max_group_num (np. 1071 to sala, nie grupa)
        for num_m in re.finditer(r'(?<!\d)(\d{1,2})(?!\d)', gm):
            val = int(num_m.group(1))
            if 1 <= val <= max_group_num:
                target_group_nums.add(val)

    if target_group_nums:
        matched_groups = [f"gr. {n}" for n in sorted(target_group_nums) if f"gr. {n}" in all_groups]
        if matched_groups:
            return matched_groups

    # Jeśli nie wykryto żadnej specyficznej grupy -> zajęcia wspólne / dla wszystkich
    return list(all_groups)


def classify_sggw_form(text: str) -> str:
    """Klasyfikuje formę zajęć na podstawie treści komórki."""
    lower = text.lower()
    if "fakultet" in lower:
        return "fakultet"
    if "ćw" in lower or "ćwiczen" in lower:
        return "cwiczenia"
    if "wykład" in lower or "wyklad" in lower or lower.strip().endswith(" w.") or " w." in lower:
        return "wyklad"
    if "lab" in lower:
        return "laboratorium"
    return "cwiczenia"


FLOOR_PATTERN = r'(?:parter|[1-4]\s*pi[eę]tr(?:o|ze)|[IVX]+\s*pi[eę]tr(?:o|ze)|[IVX]+\s*p\b|[1-4]\s*p\b)'

ROOM_PATTERNS = [
    r'(\bs\.?\s*\d{2,4}\s*b\.?\s*\d+\s+i\s+(?:s\.?\s*)?\d{1,4}\s*b\.?\s*\d+)',
    rf'(\b\d{{2,4}}\s*/\s*Aula\s*\d*(?:\s*b\.?\s*\d+)?(?:\s*(?:/|\bw\b)?\s*{FLOOR_PATTERN})?)',
    r'(Wolica\s*sala\s*\d+,\s*aula\s*\d+)',
    r'(Wolica\s*sala\s*\d+)',
    r'(sala\s*\d+,\s*aula\s*\d+)',
    r'(sala\s*109/ambulatorium/Obory)',
    r'(sala109/ambulatorium/Obory)',
    r'((?:sala\s*\d+\s*)?(?:b\.?\s*\d+\s*)?/(?:ambulatoria|ambulatorium|blok))',
    r'(06/ambul(?:\s*b\.?\s*\d+)?)',
    r'(Sala\s*\d+\s*i\s*(?:ambulatoria|ambulatorium|amb\.?)(?:\s*b\.?\s*\d+)?)',
    r'(sale\s*zakładu(?:\s*b\.?\s*\d+)?)',
    r'(sala\s*sekcyjna(?:\s*/\s*sala\s*ćwiczeniowa)?)',
    rf'((?:prosektorium|prosekt)\s*(?:w\s*)?(?:b\.?\s*\d+)?(?:\s*(?:/|\bw\b)?\s*{FLOOR_PATTERN})?)',
    rf'((?:Audytorium|Audyt\b|Aud\.?\b)\s*\d*(?:\s*(?:w\s*)?b\.?\s*\d+)?(?:\s*(?:/|\bw\b)?\s*{FLOOR_PATTERN})?)',
    rf'(Aula\s*\d*(?:\s*(?:w\s*)?b\.?\s*\d+)?(?:\s*(?:/|\bw\b)?\s*{FLOOR_PATTERN})?)',
    rf'((?:b\.?\s*\d+\s+)?(?:sala|s\.?)\s*\d+(?:[a-z])?(?:/\d+)?(?:\s*(?:w\s*)?b\.?\s*\d+)?(?:\s*(?:/|\bw\b)?\s*{FLOOR_PATTERN})?)',
    rf'((?:s\.?\s*)?[A-Z]?\d{{2,4}}\s*/\s*[A-Z]?\d{{2,4}}(?:\s*(?:w\s*)?b\.?\s*\d+)?(?:\s*(?:/|\bw\b)?\s*{FLOOR_PATTERN})?)',
    r'(\b\d{2,4}\s*w\s*b\.?\s*\d+)',
    r'(\b\d{2,4}\s*w\s*\d+)',
    r'(\bs\s*\d+\s*b\.?\s*\d+)',
    r'(\b\d{2,4}\s*b\.?\s*\d+)',
    r'(Wolica)',
    r'(Obory)',
    rf'(b\.?\s*2[2-4](?:\s*(?:/|\bw\b)?\s*{FLOOR_PATTERN})?)'
]

SUBJECT_METADATA_PATTERNS = [
    r'\b\d{1,2}[./]\d{2}[./]\d{4}\b',
    r'\bgodzina\s*\d{1,2}[:.]\d{2}(?:\s*[-–]\s*\d{1,2}[:.]\d{2})?\b',
    r'\(obecność\s*obowiązkowa[^\)]*\)',
    r'(\b\d{1,2}[:.]\d{2}\b|\bod\s*\d{1,2}[:.]\d{2}!?\b|\bdo\s*\d{1,2}[:.]\d{2}!?\b)',
    r'(od\s*\d+\s*tygodnia(?:\s*zaj[eę]ć)?|\b\d+[-–]\d+\s*tydz(?:ień)?(?:\s*zaj[eę]ć)?|(?:1|2|I|II)[.\s]*poł(?:\.|owa)?\s*sem(?:estru)?|\b1/2\b|\b2/2\b|\b\d+\s*tyg(?:\s*x\s*\d+\s*h)?\b|od\s*\d+/\d+\.[A-Z]+\.\d{2})',
    r'\(\*?\)\s*|\b(?:gr\.?|grupa|grupy|grup)\s*(?:[\d\s,–\-/]|i|\b-\b)+|\bgr\.?\s*[a-i]\b|(?<=\w)\.\d+(?:\s*,\s*\d+)+',
    r'\b(?:s\.?\s*)?\d{1,4}(?:\s*b\.?\s*\d+)?\s+i\s+(?:s\.?\s*)?\d{1,4}(?:\s*b\.?\s*\d+)?\b',
    r'\b\d{2,4}\s*/\s*Aula\s*\d*(?:\s*(?:w\s*)?b\.?\s*\d+)?',
    r'(?:Audytorium|Audyt\b|Aud\.?\b)\s*\d*(?:\s*(?:w\s*)?b\.?\s*\d+)?',
    r'Aula\s*\d*(?:\s*(?:w\s*)?b\.?\s*\d+)?',
    r'sala\s*sekcyjna(?:\s*/\s*sala\s*ćwiczeniowa)?|sale\s*zakładu(?:\s*b\.?\s*\d+)?|(?:b\.?\s*\d+\s+)?(?:sala|s\.?)\s*\d+[a-z]?(?:/\d+)?(?:\s*(?:w\s*)?b\.?\s*\d+)?',
    r'(?:s\.?\s*)?[A-Z]?\d{2,4}\s*/\s*[A-Z]?\d{2,4}(?:\s*(?:w\s*)?b\.?\s*\d+)?',
    r'\bambulatori(?:a|um)\b|\bambul\b|\bprosekt(?:orium)?\b|\bblok\b',
    r'\b\d{2,4}\s*w\s*b\.?\s*\d+|\b\d{2,4}\s*w\s*\d+|\bs\s*\d+\s*b\.?\s*\d+|\b06/ambul(?:\s*b\.?\s*\d+)?\b|\b\d{2,4}\s*b\.?\s*\d+',
    r'(?:w\s*)?budynku\s*\d+|b\.?\s*\d+',
    r'\bWolica\b|\bObory\b',
    rf'(?:/|\bw|\bna)?\s*{FLOOR_PATTERN}',
    r'\b(?:e-learning|elearning|online|zdalnie|stacjonarnie)\b',
    r'Studium\s*Języków\s*Obcych',
    r'Realizacja\s*przedmiotu\s*zgodnie\s*z\s*grafikiem[^\.]*\.?',
    r'zajęcia\s*terenowe(?:\s*\(Wolica/Obory\))?',
    r'\bćw\.?|\bćwiczenia|\bwykład\b|\bwyklad\b|\bwyk\b|\bw\.(?!\w)|\bw\b|\bfakultet\b(?!\s*\d)|\bfakultety\b'
]


def extract_room_and_subject(text: str) -> Tuple[str, str, str]:
    """
    Wyodrębnia nazwę przedmiotu, salę oraz uwagi z tekstu komórki.
    """
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    if not lines:
        return ("", "", "")

    full_text = " ".join(lines)
    is_elearning = bool(re.search(r'\b(?:e-learning|elearning|online|zdalnie)\b', full_text, re.IGNORECASE))

    # 1. Wykrywanie sali
    room = ""
    for rp in ROOM_PATTERNS:
        m = re.search(rp, full_text, re.IGNORECASE)
        if m:
            room = m.group(1).strip()
            break

    if not room and is_elearning:
        room = "online"

    # 2. Wyodrębnienie i oczyszczenie nazwy przedmiotu
    first_line = lines[0]
    # Obsługa przypadku "Fakultety" w pierwszej linii, gdy faktyczna nazwa jest w drugiej
    if first_line.lower().strip() in ("fakultet", "fakultety") and len(lines) > 1:
        first_line = lines[1]

    subject = first_line
    for pat in SUBJECT_METADATA_PATTERNS:
        subject = re.sub(pat, '', subject, flags=re.IGNORECASE)
    subject = re.sub(r'[*!,;:/()]+', ' ', subject)
    subject = ' '.join(subject.split()).strip()

    # Usunięcie końcowej cyfry (o ile to nie jest np. Fakultet 1 / Fakultet 2)
    if not re.search(r'\bfakultet\s+\d+$', subject, re.IGNORECASE):
        subject = re.sub(r'\s+[1-9]$', '', subject).strip()
    subject = re.sub(r'^[.\-\s–—]+|[.\-\s–—]+$', '', subject)
    subject = re.sub(r'\s+\bi\b$', '', subject).strip()
    subject = re.sub(r'^[.\-\s–—]+|[.\-\s–—]+$', '', subject)

    # Zachowanie kropki dla skrótów typu "odp psych."
    if subject.endswith("odp psych"):
        subject = subject + "."

    # Sprawdzenie kanonicznej mapy przedmiotów
    canon = CANONICAL_SUBJECT_MAP.get(subject.lower())
    if canon:
        subject = canon

    if not subject:
        # Fallback do pierwszej linii bez dopisku formy
        subject = re.sub(r'\s*(?:wykład|wyklad|ćw\.?|cw\.?|ćwiczenia|fakultet).*$', '', first_line, flags=re.IGNORECASE).strip()
        if not subject:
            subject = first_line

    # 3. Uwagi / tygodnie
    notes = ""
    week_notes = re.findall(
        r'(?:od\s*\d+\s*tygodnia|1-\d+\s*tydzień|(?:1|2|I|II)[.\s]*poł(?:\.|owa)?\s*sem(?:estru)?|od\s*\d+:\d+!|\b\d+[-–]\d+\s*tydz\b)',
        full_text,
        re.IGNORECASE
    )
    if week_notes:
        notes = ", ".join(week_notes)

    return (subject, room, notes)


def parse_sggw_rotations_excel(wb_or_sheet: Any) -> Dict[str, Dict[str, Dict[Any, Any]]]:
    """
    Parsuje arkusz rotacji klinicznych (semestr 11 VI rok WMW).
    Zwraca strukturę grup: 14 stażowych (gr. 1a-7b) oraz 7 ogólnych dziekańskich (gr. 1-7).
    """
    if isinstance(wb_or_sheet, str):
        wb = openpyxl.load_workbook(wb_or_sheet, data_only=True)
        sheet = wb.active
    elif hasattr(wb_or_sheet, "active"):
        sheet = wb_or_sheet.active
    else:
        sheet = wb_or_sheet

    clinics: Dict[int, str] = {}
    for col in range(3, min(20, sheet.max_column + 1)):
        v5 = sheet.cell(5, col).value
        v6 = sheet.cell(6, col).value
        parts = []
        if v5:
            parts.append(str(v5).strip())
        if v6 and str(v6).strip() not in parts:
            parts.append(str(v6).strip())
        name = " ".join(parts).strip()
        if name:
            clinics[col] = name

    subgroups = [f"gr. {i}{letter}" for i in range(1, 8) for letter in ('a', 'b')]
    main_groups = [f"gr. {i}" for i in range(1, 8)]
    all_groups = subgroups + main_groups

    result: Dict[str, Dict[str, Dict[Any, Any]]] = {
        g: {d: {} for d in ["PON", "WT", "ŚR", "CZW", "PT", "SOB", "ND"]} for g in all_groups
    }

    def _parse_range(text: str) -> Tuple[Optional[datetime.date], Optional[datetime.date]]:
        m = re.search(r'(\d{1,2})\.(\d{2})\s*[-–]\s*(\d{1,2})\.(\d{2})', text)
        if not m:
            return None, None
        d1, m1, d2, m2 = int(m.group(1)), int(m.group(2)), int(m.group(3)), int(m.group(4))
        y1 = 2026 if m1 >= 10 else 2027
        y2 = 2026 if m2 >= 10 else 2027
        return datetime(y1, m1, d1).date(), datetime(y2, m2, d2).date()

    weekday_map = {0: "PON", 1: "WT", 2: "ŚR", 3: "CZW", 4: "PT", 6: "ND"}

    # Wiersz 7: Wykłady ZZS dla wszystkich grup
    r7_c1 = str(sheet.cell(7, 1).value or "")
    d1_r7, d2_r7 = _parse_range(r7_c1)
    if d1_r7 and d2_r7:
        cur = d1_r7
        while cur <= d2_r7:
            if cur.weekday() in weekday_map:
                day_code = weekday_map[cur.weekday()]
                for g in all_groups:
                    slot_key = "wyklad_zzs"
                    result[g][day_code][slot_key] = {
                        "subject": "Wykłady ZZS",
                        "raw_subject": "Wykłady ZZS",
                        "teacher": "",
                        "hours": "08:00 - 14:30",
                        "room": "CMT",
                        "height": 26,
                        "colspan": 1,
                        "start_date": d1_r7.isoformat(),
                        "weeks": 1,
                        "interval_weeks": 1,
                        "start_week": 1,
                        "half_semester": None,
                        "form": "wyklad",
                        "notes": "",
                        "groups": list(all_groups)
                    }
            cur += timedelta(days=1)

    # Wiersze 8-25: Rotacje kliniczne
    for r in range(8, sheet.max_row + 1):
        c1_val = str(sheet.cell(r, 1).value or "")
        d1, d2 = _parse_range(c1_val)
        if not d1 or not d2:
            continue

        week_days = []
        cur = d1
        while cur <= d2:
            if cur.weekday() in weekday_map:
                week_days.append((cur, weekday_map[cur.weekday()]))
            cur += timedelta(days=1)

        for col, clinic_name in clinics.items():
            cell_val = sheet.cell(r, col).value
            if not cell_val:
                continue
            text = str(cell_val).strip()
            if text.lower() in ('x', '-', ''):
                continue

            found_subs = re.findall(r'(\d+)\s*([ab])', text, re.IGNORECASE)
            if not found_subs:
                continue

            notes_match = re.search(r'\(([^)]*(?:rozr[óo]d|CMT|staż)[^)]*)\)', text, re.IGNORECASE)
            notes = notes_match.group(1).strip() if notes_match else ""

            target_groups = set()
            for num, letter in found_subs:
                sub = f"gr. {num}{letter.lower()}"
                main = f"gr. {num}"
                target_groups.add(sub)
                target_groups.add(main)

            sorted_groups = sorted(list(target_groups))
            for d_date, day_code in week_days:
                slot_key = f"staz_{r}_{col}"
                subject_name = f"Staż - {clinic_name}"
                lesson_entry = {
                    "subject": subject_name,
                    "raw_subject": clinic_name,
                    "teacher": "",
                    "hours": "08:00 - 14:00",
                    "room": clinic_name,
                    "height": 24,
                    "colspan": 1,
                    "start_date": d1.isoformat(),
                    "weeks": 1,
                    "interval_weeks": 1,
                    "start_week": 1,
                    "half_semester": None,
                    "form": "staz",
                    "notes": notes,
                    "groups": sorted_groups
                }
                for g in target_groups:
                    result[g][day_code][slot_key] = lesson_entry

    return result


def parse_sggw_excel(wb_or_path: Union[str, openpyxl.Workbook]) -> Dict[str, Dict[str, Dict[Any, Any]]]:
    """
    Główna funkcja parsująca arkusz planu zajęć SGGW.
    Zwraca słownik grup i ich harmonogramów.
    """
    if isinstance(wb_or_path, str):
        wb = openpyxl.load_workbook(wb_or_path, data_only=True)
    else:
        wb = wb_or_path

    sheet = wb.active

    # Sprawdzenie czy arkusz to rotacje kliniczne (np. semestr 11)
    title_val = str(sheet.cell(1, 1).value or sheet.cell(1, 2).value or sheet.cell(1, 3).value or "")
    if "STAŻE" in title_val.upper() or "STAZE" in title_val.upper() or "ROTACJ" in title_val.upper():
        return parse_sggw_rotations_excel(sheet)

    # 1. Odczyt kolumn czasowych (szukamy w wierszach 1-3)
    time_columns: Dict[int, Tuple[str, str]] = {}
    time_row = 2
    for r_cand in [2, 1, 3]:
        for col in range(1, sheet.max_column + 1):
            v = sheet.cell(r_cand, col).value
            if v and re.search(r'\d{1,2}[:.]\d{2}\s*[-–]\s*\d{1,2}[:.]\d{2}', str(v)):
                time_row = r_cand
                break
        if any(re.search(r'\d{1,2}[:.]\d{2}\s*[-–]\s*\d{1,2}[:.]\d{2}', str(sheet.cell(time_row, c).value or "")) for c in range(1, min(10, sheet.max_column + 1))):
            break

    for col in range(1, sheet.max_column + 1):
        v = sheet.cell(time_row, col).value
        if v:
            v_str = str(v).strip()
            m = re.search(r'(\d{1,2}[:.]\d{2})\s*[-–]\s*(\d{1,2}[:.]\d{2})', v_str)
            if m:
                sh = _normalize_time_24h(m.group(1))
                eh = _normalize_time_24h(m.group(2))
                time_columns[col] = (sh, eh)

    if not time_columns:
        # Fallback dla arkuszy rotacji
        return parse_sggw_rotations_excel(sheet)

    # 2. Wykrycie liczby grup (z wiersza 1 lub domyślnie 7)
    title_val = str(sheet.cell(1, 1).value or sheet.cell(1, 2).value or sheet.cell(1, 3).value or "")
    m_groups = re.search(r'(\d+)\s*grup', title_val, re.IGNORECASE)
    total_groups_num = int(m_groups.group(1)) if m_groups else 7
    all_groups = [f"gr. {i}" for i in range(1, total_groups_num + 1)]

    # Przygotowanie struktury wyjściowej dla każdej grupy
    result: Dict[str, Dict[str, Dict[int, Any]]] = {
        g: {d: {} for d in ["PON", "WT", "ŚR", "CZW", "PT"]} for g in all_groups
    }

    # 3. Indeksowanie scalonych komórek
    merged_map: Dict[Tuple[int, int], Tuple[int, int, int, int]] = {}
    for rng in sheet.merged_cells.ranges:
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                merged_map[(r, c)] = (rng.min_row, rng.min_col, rng.max_row, rng.max_col)

    # 4. Przejście po wierszach i dniach tygodnia
    current_day = ""
    processed_top_lefts = set()
    start_c = min(time_columns.keys()) if time_columns else 2

    for r in range(time_row + 1, sheet.max_row + 1):
        # Sprawdzenie dnia w kolumnie 1 lub 2
        day_cell = sheet.cell(r, 1).value or sheet.cell(r, 2).value
        if day_cell:
            day_text = str(day_cell).strip().lower().split("\n")[0].rstrip(".:")
            for k, code in DAY_NAME_MAP.items():
                if day_text == k or day_text.startswith(k):
                    current_day = code
                    break

        if not current_day:
            continue

        for c in range(start_c, sheet.max_column + 1):
            cell_val = sheet.cell(r, c).value
            if not cell_val or not str(cell_val).strip():
                continue

            content = str(cell_val).strip()
            # Ignorowanie legendy / przypisów na dole
            if (content.startswith("* Szczegółowy") or
                content.startswith("Plan zajęć") or
                re.search(r'^(?:\*\s*)?(?:Wybór\s*fakultetu|\*\s*Fakultet|Fakultet\s*\d+:)', content, re.IGNORECASE)):
                continue

            # Sprawdzenie czy komórka jest częścią zakresu scalonego
            if (r, c) in merged_map:
                min_r, min_c, max_r, max_c = merged_map[(r, c)]
                if (min_r, min_c) in processed_top_lefts:
                    continue
                processed_top_lefts.add((min_r, min_c))
                start_col = min_c
                end_col = max_c
            else:
                start_col = c
                end_col = c

            # Wyznaczenie domyślnych godzin z kolumn
            def_start = time_columns.get(start_col, ("08:00", "09:00"))[0]
            def_end = time_columns.get(end_col, ("09:00", "10:00"))[1]

            actual_start, actual_end = parse_time_bounds(content, def_start, def_end)
            hours_str = f"{actual_start} - {actual_end}"

            # Slot bazowy liczony jako minuty od 07:00 / 5
            try:
                sh_val, sm_val = map(int, actual_start.split(":"))
                eh_val, em_val = map(int, actual_end.split(":"))
                base_slot = ((sh_val - 7) * 60 + sm_val) // 5
                duration_min = (eh_val * 60 + em_val) - (sh_val * 60 + sm_val)
                height = duration_min // 5
            except Exception:
                base_slot = (start_col - 3) * 12 + 12
                height = (end_col - start_col + 1) * 12

            # Wyodrębnienie grup, przedmiotu i sali
            target_groups = extract_groups_from_text(content, all_groups)
            form_code = classify_sggw_form(content)
            subject, room, notes = extract_room_and_subject(content)

            # Sprawdzenie czy komórka łączy dwa przedmioty w 1. i 2. połowie semestru
            m_double_half = re.search(
                r'^(.*?)(?:1[.\s]*poł|I\s*poł)[.\s]*sem(?:estru)?\.?\s*(.*?)(?:2[.\s]*poł|II\s*poł)[.\s]*sem(?:estru)?\.?(.*)$',
                content,
                re.IGNORECASE | re.DOTALL
            )

            lessons_to_add: List[Dict[str, Any]] = []

            if m_double_half:
                part1_raw = m_double_half.group(1).strip()
                part2_raw = m_double_half.group(2).strip()
                extra_raw = m_double_half.group(3).strip()
                s1, _, _ = extract_room_and_subject(part1_raw + " " + extra_raw)
                s2, _, _ = extract_room_and_subject(part2_raw + " " + extra_raw)

                lessons_to_add.append({
                    "subject": s1,
                    "raw_subject": s1,
                    "teacher": "",
                    "hours": hours_str,
                    "room": room or "brak sali",
                    "height": height,
                    "colspan": 1,
                    "start_date": "2026-10-01",
                    "weeks": 7,
                    "interval_weeks": 1,
                    "start_week": 1,
                    "half_semester": 1,
                    "form": form_code,
                    "notes": notes,
                    "groups": list(target_groups)
                })
                lessons_to_add.append({
                    "subject": s2,
                    "raw_subject": s2,
                    "teacher": "",
                    "hours": hours_str,
                    "room": room or "brak sali",
                    "height": height,
                    "colspan": 1,
                    "start_date": "2026-11-23",
                    "weeks": 8,
                    "interval_weeks": 1,
                    "start_week": 1,
                    "half_semester": 2,
                    "form": form_code,
                    "notes": notes,
                    "groups": list(target_groups)
                })
            else:
                # Tygodnie i cykle
                weeks = 15
                interval_weeks = 1
                start_week = 1
                half_semester = None
                start_date = "2026-10-01"

                m_week_range = re.search(r'1\s*-\s*(\d+)\s*tydz', content, re.IGNORECASE)
                m_from_week = re.search(r'od\s*(\d+)\s*tygodnia', content, re.IGNORECASE)

                if m_week_range:
                    weeks = int(m_week_range.group(1))
                    start_date = "2026-10-01"
                elif m_from_week:
                    start_week = int(m_from_week.group(1))
                    weeks = max(1, 15 - start_week + 1)
                    # Obliczenie daty startu: poniedziałek 1. tygodnia to 2026-10-05
                    base_mon = datetime(2026, 10, 5)
                    start_dt = base_mon + timedelta(days=(start_week - 1) * 7)
                    start_date = start_dt.strftime("%Y-%m-%d")
                elif re.search(r'1[.\s]*poł(?:\.|owa)?\s*sem', content, re.IGNORECASE):
                    weeks = 7
                    half_semester = 1
                    start_date = "2026-10-01"
                elif re.search(r'2[.\s]*poł(?:\.|owa)?\s*sem', content, re.IGNORECASE):
                    weeks = 8
                    half_semester = 2
                    start_date = "2026-11-23"

                lessons_to_add.append({
                    "subject": subject,
                    "raw_subject": subject,
                    "teacher": "",
                    "hours": hours_str,
                    "room": room or "brak sali",
                    "height": height,
                    "colspan": 1,
                    "start_date": start_date,
                    "weeks": weeks,
                    "interval_weeks": interval_weeks,
                    "start_week": start_week,
                    "half_semester": half_semester,
                    "form": form_code,
                    "notes": notes,
                    "groups": list(target_groups)
                })

            # Zapisanie do każdej zainteresowanej grupy
            for grp in target_groups:
                day_slots = result[grp][current_day]
                for l_dict in lessons_to_add:
                    # Obsługa ewentualnych kolizji slotów (np. 2 zajęcia o tej samej godzinie)
                    if base_slot not in day_slots:
                        day_slots[base_slot] = l_dict
                    else:
                        existing = day_slots[base_slot]
                        if isinstance(existing, dict) and "subject" in existing:
                            day_slots[base_slot] = {0: existing, 1: l_dict}
                        elif isinstance(existing, dict):
                            existing[len(existing)] = l_dict

    return result
